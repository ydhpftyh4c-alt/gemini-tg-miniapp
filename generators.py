import os
import re
import json
import urllib.parse
from io import BytesIO
from typing import List, Dict, Any, Optional

import httpx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

TEMP_DIR = os.path.join(os.path.dirname(__file__), "generated_files")
os.makedirs(TEMP_DIR, exist_ok=True)

# Register Arial or system font for Cyrillic PDF support
CYRILLIC_FONT = "Helvetica"
try:
    # Try finding Arial on Windows
    win_font = "C:\\Windows\\Fonts\\arial.ttf"
    if os.path.exists(win_font):
        pdfmetrics.registerFont(TTFont("Arial", win_font))
        CYRILLIC_FONT = "Arial"
except Exception:
    pass

# --- 1. Image Generation (Flux / Pollinations with Gemini prompt enhancement) ---
async def generate_image_bytes(user_prompt: str, gemini_caller) -> tuple[bytes, str]:
    """
    Enhances prompt using Gemini, then fetches high-res image.
    """
    try:
        # Prompt enhancement
        enhanced = await gemini_caller(
            user_id=0,
            user_message=f"Translate and expand this into a detailed English prompt for an image generator (Flux/SDXL). Output ONLY the English prompt, nothing else: {user_prompt}",
            model_name="gemini-3.8-flash"
        )
        clean_prompt = enhanced.strip().strip('"').strip("'")
    except Exception:
        clean_prompt = user_prompt

    encoded = urllib.parse.quote(clean_prompt)
    url = f"https://image.pollinations.ai/prompt/{encoded}?width=1024&height=1024&nologo=true&seed={abs(hash(clean_prompt)) % 100000}"
    
    async with httpx.AsyncClient(timeout=30.0) as client:
        resp = await client.get(url)
        if resp.status_code == 200:
            return resp.content, clean_prompt
        raise Exception(f"Image generation failed with HTTP {resp.status_code}")

# --- 2. DOCX Word Document Generation ---
async def generate_docx_file(topic: str, gemini_caller) -> str:
    """
    Generates a structured Word document (.docx) using Gemini.
    """
    prompt = (
        f"Напиши подробный, профессиональный структурированный отчёт/документ на тему: '{topic}'.\n"
        "Формат ответа:\n"
        "Раздели документ на чёткие секции. Заголовки пиши как '## Заголовок', "
        "подзаголовки как '### Подзаголовок', списки через '-', текст абзацами. "
        "Сделай документ качественным, развёрнутым и полезным."
    )
    
    content = await gemini_caller(
        user_id=0,
        user_message=prompt,
        model_name="gemini-3.8-flash"
    )
    
    doc = Document()
    
    # Title
    title = doc.add_heading(level=0)
    title_run = title.add_run(topic.title())
    title_run.font.color.rgb = RGBColor(0x2A, 0x2A, 0x72)
    title_run.font.bold = True
    
    # Body parser
    for line in content.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("## "):
            h = doc.add_heading(line[3:], level=1)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x3B, 0x49, 0xDF)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif re.match(r"^\d+\.\s", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", line), style='List Number')
        else:
            doc.add_paragraph(line)
            
    filename = f"report_{abs(hash(topic)) % 100000}.docx"
    filepath = os.path.join(TEMP_DIR, filename)
    doc.save(filepath)
    return filepath

# --- 3. PDF Document Generation ---
async def generate_pdf_file(topic: str, gemini_caller) -> str:
    """
    Generates a PDF document using ReportLab and Gemini.
    """
    prompt = (
        f"Напиши структурированный отчёт на тему: '{topic}'. "
        "Включи введение, 3-4 ключевых раздела с выводами и заключение. Пиши обычным текстом с заголовками."
    )
    
    content = await gemini_caller(
        user_id=0,
        user_message=prompt,
        model_name="gemini-3.8-flash"
    )
    
    filename = f"document_{abs(hash(topic)) % 100000}.pdf"
    filepath = os.path.join(TEMP_DIR, filename)
    
    doc = SimpleDocTemplate(filepath, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName=CYRILLIC_FONT,
        fontSize=20,
        leading=24,
        textColor='#1a1a4b',
        spaceAfter=14
    )
    
    heading_style = ParagraphStyle(
        'DocHeading',
        parent=styles['Heading2'],
        fontName=CYRILLIC_FONT,
        fontSize=14,
        leading=18,
        textColor='#3344aa',
        spaceBefore=10,
        spaceAfter=6
    )
    
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['BodyText'],
        fontName=CYRILLIC_FONT,
        fontSize=10,
        leading=14,
        spaceAfter=6
    )
    
    story = []
    story.append(Paragraph(topic.title(), title_style))
    story.append(Spacer(1, 12))
    
    for line in content.split("\n"):
        line = line.strip()
        if not line:
            continue
        line_clean = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if line.startswith("#"):
            text = re.sub(r"^#+\s*", "", line_clean)
            story.append(Paragraph(text, heading_style))
        else:
            story.append(Paragraph(line_clean, body_style))
            
    doc.build(story)
    return filepath

# --- 4. PPTX PowerPoint Presentation Generation ---
async def generate_pptx_file(topic: str, gemini_caller) -> str:
    """
    Generates a PowerPoint presentation (.pptx) with slides from Gemini JSON structure.
    """
    prompt = (
        f"Создай презентацию на тему: '{topic}'.\n"
        "Сформируй ровно 5 слайдов в формате чистого JSON списка без лишних слов:\n"
        "[\n"
        "  {\"title\": \"Заголовок слайда 1\", \"bullets\": [\"Пункт 1\", \"Пункт 2\", \"Пункт 3\"]},\n"
        "  {\"title\": \"Заголовок слайда 2\", \"bullets\": [\"Пункт 1\", \"Пункт 2\", \"Пункт 3\"]}\n"
        "]"
    )
    
    json_text = await gemini_caller(
        user_id=0,
        user_message=prompt,
        model_name="gemini-3.8-flash"
    )
    
    # Extract JSON array
    slides_data = []
    try:
        match = re.search(r"\[[\s\S]*\]", json_text)
        if match:
            slides_data = json.loads(match.group(0))
    except Exception:
        pass
        
    if not slides_data:
        slides_data = [
            {"title": topic.title(), "bullets": ["Введение в тему", "Ключевые аспекты", "Практическое применение"]},
            {"title": "Основные факты", "bullets": ["Фактор 1", "Фактор 2", "Фактор 3"]},
            {"title": "Заключение", "bullets": ["Итоги", "Выводы", "Вопросы"]}
        ]
        
    prs = Presentation()
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    
    for i, slide_info in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)
        
        # Slide Background (Modern dark slate / navy)
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = PptxRGBColor(0x16, 0x1B, 0x2E)
        
        # Title box
        txBox = slide.shapes.add_textbox(PptxInches(1.0), PptxInches(0.8), PptxInches(11.3), PptxInches(1.2))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = slide_info.get("title", f"Слайд {i+1}")
        p.font.size = PptxPt(36)
        p.font.bold = True
        p.font.color.rgb = PptxRGBColor(0x6C, 0x63, 0xFF)
        
        # Bullets box
        bullets = slide_info.get("bullets", [])
        txBox2 = slide.shapes.add_textbox(PptxInches(1.0), PptxInches(2.3), PptxInches(11.3), PptxInches(4.5))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        
        for j, bullet in enumerate(bullets):
            p2 = tf2.add_paragraph() if j > 0 else tf2.paragraphs[0]
            p2.text = f"•  {bullet}"
            p2.font.size = PptxPt(22)
            p2.font.color.rgb = PptxRGBColor(0xE0, 0xE0, 0xEE)
            p2.space_after = PptxPt(16)
            
    filename = f"presentation_{abs(hash(topic)) % 100000}.pptx"
    filepath = os.path.join(TEMP_DIR, filename)
    prs.save(filepath)
    return filepath
